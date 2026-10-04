#!/usr/bin/env python3
"""Build the five-slide Shortly DevOps project presentation."""

from __future__ import annotations

import importlib.util
import os
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]


def ensure_pptx_runtime() -> None:
    if importlib.util.find_spec("pptx"):
        return
    runtime = (
        Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies/python/bin/python3"
    )
    if runtime.exists() and str(runtime) != sys.executable:
        os.execv(str(runtime), [str(runtime), str(Path(__file__).resolve()), *sys.argv[1:]])
    raise SystemExit("python-pptx is required; install requirements-report.txt and retry.")


ensure_pptx_runtime()

from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

OUT = ROOT / "report"
SHOT = ROOT / "docs" / "screenshots"
SLIDE_W, SLIDE_H = 13.333, 7.5
NAVY = "172B3A"
INK = "18313F"
TEAL = "1A8C8C"
PALE = "EEF4F5"
MID = "D2E2E5"
GRAY = "60727C"
WHITE = "FFFFFF"
AMBER = "C07A1A"
FONT = "Arial"
MEMBERS = [
    ("Harsh Ledwani", "23070122100"),
    ("Shivam Kapure", "23070122113"),
    ("Kashyup Gaud", "23070122114"),
    ("Kisna Kanti", "23070122116"),
]


def color(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def textbox(slide, text, x, y, w, h, size=14, fill=INK, bold=False, align=None, margin=0.04):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color(fill)
    return shape


def box(slide, text, x, y, w, h, fill=PALE, line=MID, size=14, bold=True, ink=INK):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    shape.line.color.rgb = color(line)
    shape.line.width = Pt(1)
    shape.text_frame.clear()
    shape.text_frame.word_wrap = True
    shape.text_frame.margin_left = shape.text_frame.margin_right = Inches(0.09)
    shape.text_frame.margin_top = shape.text_frame.margin_bottom = Inches(0.04)
    shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = shape.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color(ink)
    return shape


def line(slide, x1, y1, x2, y2, ink=TEAL, width=1.5, dashed=False):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = color(ink)
    shape.line.width = Pt(width)
    shape.line.end_arrowhead = True
    if dashed:
        shape.line.dash_style = 2
    return shape


def header(slide, title, subtitle, page):
    textbox(slide, title, 0.52, 0.25, 12.25, 0.55, 27, NAVY, True)
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(0.91), Inches(0.78), Inches(0.055)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = color(TEAL)
    bar.line.fill.background()
    textbox(slide, subtitle, 0.54, 0.99, 12.2, 0.34, 12.5, GRAY)
    textbox(slide, f"SHORTLY  /  DEVOPS CA-2                                      {page:02d}",
            0.55, 7.13, 12.1, 0.2, 8.5, GRAY)


def add_image(slide, filename, x, y, w, h, caption=None):
    path = SHOT / filename
    if not path.is_file():
        textbox(slide, f"Image missing: {filename}", x, y, w, h, 12, AMBER, True, PP_ALIGN.CENTER)
        return
    with Image.open(path) as source:
        image_ratio = source.width / source.height
    frame_ratio = w / h
    if image_ratio > frame_ratio:
        image_w, image_h = w, w / image_ratio
    else:
        image_h, image_w = h, h * image_ratio
    slide.shapes.add_picture(
        str(path), Inches(x + (w - image_w) / 2), Inches(y + (h - image_h) / 2),
        width=Inches(image_w), height=Inches(image_h)
    )
    if caption:
        textbox(slide, caption, x, y + h + 0.015, w, 0.25, 10, GRAY, True, PP_ALIGN.CENTER)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def slide_architecture(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "Shortly turns long URLs into shared, trackable links",
           "Architecture and project team", 1)
    # This is the team's captured architecture diagram, preserved as project evidence.
    add_image(slide, "27-pipeline-diagram.png", 0.62, 1.43, 12.05, 2.94,
              "Captured project architecture diagram")
    add_image(slide, "APP_frontend.png", 0.78, 4.72, 3.46, 2.25,
              "The app creates and lists short links")
    textbox(slide, "PROJECT GROUP", 4.72, 4.64, 6.9, 0.28, 11, TEAL, True)
    for i, (name, roll) in enumerate(MEMBERS):
        y = 4.98 + i * 0.44
        textbox(slide, name, 4.74, y, 3.25, 0.32, 14, INK, True)
        textbox(slide, f"Roll no.  {roll}", 8.12, y, 3.4, 0.32, 13, GRAY)
    notes(slide, "Shortly is a FastAPI URL shortener. The browser reaches the API through nginx Ingress and a Kubernetes Service. Two or more stateless API pods share link state through Redis. Redis writes to a persistent volume. Prometheus scrapes application metrics, while Grafana and Alertmanager show and evaluate them. GitHub Actions builds and checks each release, and a self-hosted runner deploys to the local Minikube cluster. The team members and roll numbers come from the project report notes.")


def slide_pipeline(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "GitHub Actions checks the image before Minikube deploys it",
           "Hosted build and scan jobs hand the verified image to a guarded local runner", 2)
    y, h, w = 1.75, 0.86, 2.05
    xs = [0.58, 3.08, 5.58, 8.08, 10.58]
    labels = ["Push to main", "Lint and tests", "Build, scan\nand SBOM", "Publish image\nto GHCR", "Verify digest\nand deploy"]
    fills = [PALE, PALE, PALE, PALE, "FFF5E6"]
    strokes = [TEAL, TEAL, TEAL, TEAL, AMBER]
    for x, label, fill, stroke in zip(xs, labels, fills, strokes):
        box(slide, label, x, y, w, h, fill, stroke, 13)
    for x in [2.65, 5.15, 7.65, 10.15]:
        line(slide, x, y + h / 2, x + 0.4, y + h / 2)
    textbox(slide, "GITHUB-HOSTED", 0.6, 1.48, 2.0, 0.2, 9.5, TEAL, True)
    textbox(slide, "SELF-HOSTED", 10.56, 1.48, 2.0, 0.2, 9.5, AMBER, True)
    box(slide, "Verify rollout, smoke\nand version", 10.58, 3.05, 2.05, 0.68,
        "E8F3F1", TEAL, 11.5)
    line(slide, 11.6, 2.61, 11.6, 3.05, TEAL)
    box(slide, "Pass: release succeeds", 10.58, 3.88, 2.05, 0.57,
        "E8F3F1", TEAL, 10.5)
    box(slide, "Failure: undo rollout; job stays red", 7.08, 3.88, 2.85, 0.57,
        "FFF5E6", AMBER, 10.5)
    line(slide, 11.6, 3.73, 11.6, 3.88, TEAL)
    line(slide, 10.58, 3.73, 8.5, 3.73, AMBER, 1.3)
    line(slide, 8.5, 3.73, 8.5, 3.88, AMBER, 1.3)
    add_image(slide, "12-deploy-job-summary.png", 0.62, 4.72, 5.05, 1.94,
              "Captured successful workflow run: all jobs passed in 1m 58s")
    textbox(slide, "The deploy job verifies the same image digest that passed the hosted scan.\n\n"
            "A verification failure triggers rollback and leaves the workflow visibly failed.",
            6.12, 5.05, 6.45, 1.12, 15, INK, True)
    notes(slide, "The workflow runs lint and tests on GitHub hosted runners, builds a container, scans it with Trivy and creates an SBOM, then publishes it to GHCR. The self-hosted job checks its cluster context and verifies the published image digest before deploying. It then waits for rollout probes and checks smoke requests and the version. The screenshot shows a successful workflow run lasting 1 minute 58 seconds. The failure branch automatically rolls back, while the workflow remains red so the release failure is visible.")


def slide_observability(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "Grafana makes traffic, errors, and latency visible",
           "Prometheus scrapes the app; the dashboard turns those metrics into signals", 3)
    add_image(slide, "01-baseline-overview.png", 0.62, 1.52, 7.5, 4.92,
              "Captured baseline dashboard: 100% availability and 0% 5xx")
    add_image(slide, "03-latency-p95-alert.png", 8.45, 1.55, 4.23, 2.12,
              "Latency scenario: p95 rises above the 500 ms threshold")
    add_image(slide, "04-errors-5xx-alert.png", 8.45, 4.02, 4.23, 2.12,
              "Error scenario: 5xx rate rises during injected failures")
    textbox(slide, "Saved baseline: 3,753 requests · 0 unexpected failures · client p95 8 ms",
            0.7, 6.73, 12.0, 0.25, 12, INK, True, PP_ALIGN.CENTER)
    notes(slide, "The dashboard presents pod count, availability, request rate, 5xx rate, p95 latency and Redis health. The baseline run recorded 3,753 requests, zero unexpected failures and an 8 ms client p95. Separate captured Grafana views show the latency fault rising above the configured threshold and the 5xx rate increasing during the error scenario. The scenario summary records that both alerts fired and later recovered. These are saved screenshots and report values, not a live rerun during deck preparation.")


def slide_challenges(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "Readiness alone did not catch a broken release",
           "The release stayed Ready while ordinary redirect requests failed", 4)
    textbox(slide, "WHAT HAPPENED", 0.7, 1.6, 4.6, 0.28, 11, TEAL, True)
    textbox(slide, "The bad image passed Kubernetes readiness, but smoke checks found redirect, URL validation, and blocked-domain requests returning 500.",
            0.7, 1.98, 5.35, 1.06, 17, INK, True)
    textbox(slide, "HOW THE PIPELINE RESPONDED", 0.7, 3.25, 5.0, 0.28, 11, AMBER, True)
    textbox(slide, "It rolled the Deployment back to revision 58. The workflow stayed failed so the team could see that verification stopped the release.",
            0.7, 3.62, 5.35, 0.97, 16, INK)
    textbox(slide, "LESSON", 0.7, 4.83, 2.0, 0.28, 11, TEAL, True)
    textbox(slide, "Health probes check process and dependency health. Smoke tests check user behavior.",
            0.7, 5.18, 5.35, 0.68, 16, INK, True)
    add_image(slide, "11-pipeline-autorollback.png", 6.45, 1.55, 4.12, 4.48,
              "Captured failed release with automatic rollback")
    add_image(slide, "16-kubectl-pods.png", 10.75, 2.35, 2.05, 0.39,
              "Ready API and Redis pods")
    textbox(slide, "Ansible runs on a supported Ubuntu 24.04 target. The first apply changed 31 items; the repeat apply changed 0.",
            6.45, 6.58, 6.25, 0.48, 11.5, GRAY)
    notes(slide, "The bad-release-errors scenario demonstrated a gap between pod readiness and correct application behavior. Readiness remained green while normal redirects and other smoke requests failed. The workflow ran its rollback helper, restored revision 58 and returned a failed status. The pods screenshot shows the stable app and Redis state after recovery. Ansible's positive results come from the disposable Ubuntu 24.04 container: the first apply changed 31 items, and the second changed zero. The Mac host was rejected by the playbook before changes.")


def slide_lessons(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "The build taught us to verify releases from the user's side",
           "Results, limits, and the next improvements", 5)
    lessons = [
        ("01", "Keep readiness and user-facing checks separate."),
        ("02", "Deploy the exact scanned image, then verify requests and version."),
        ("03", "Use observed resource needs and keep alerts tied to real traffic."),
        ("04", "Run configuration management on a supported, repeatable host."),
    ]
    for i, (num, lesson) in enumerate(lessons):
        y = 1.62 + i * 0.58
        textbox(slide, num, 0.74, y, 0.48, 0.38, 13, TEAL, True)
        textbox(slide, lesson, 1.25, y, 6.4, 0.38, 15, INK, i == 0)
    box(slide, "8 of 9 scenarios passed in the saved full run", 8.05, 1.67, 4.55, 0.7,
        PALE, TEAL, 14)
    textbox(slide, "The focused abuse rerun later passed. The full suite was not rerun, so the original 8/9 result remains the reported full-run outcome.",
            8.13, 2.56, 4.37, 0.75, 12.5, GRAY)
    add_image(slide, "14-runner-online.png", 1.35, 4.28, 3.55, 2.25,
              "Self-hosted deployment runner")
    add_image(slide, "15-environments.png", 7.55, 4.28, 3.55, 2.25,
              "GitHub environment configured for local Minikube")
    textbox(slide, "Next: repeat on a multi-node cluster, add a real alert receiver, and isolate the deployment runner.",
            0.78, 6.78, 11.9, 0.23, 11, GRAY, True, PP_ALIGN.CENTER)
    notes(slide, "The main learning is that a healthy process does not guarantee a healthy user request. The pipeline now verifies the exact image after deploy and automatically undoes a release when checks fail. Observed load results and alert behavior informed resource and monitoring choices. The saved full scenario run passed eight of nine tests; a focused abuse rerun passed afterward, but the full suite was not rerun. The project uses a single-node local Minikube cluster, one Redis replica, a null Alertmanager receiver, and a local self-hosted runner. Future work is to repeat this on a multi-node cluster, route alerts to a real receiver and isolate the runner.")


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    prs.core_properties.title = "Shortly DevOps CA-2"
    prs.core_properties.subject = "Architecture, deployment pipeline, monitoring, challenges, and lessons learned"
    prs.core_properties.author = ", ".join(name for name, _ in MEMBERS)
    for builder in [slide_architecture, slide_pipeline, slide_observability, slide_challenges, slide_lessons]:
        builder(prs)
    path = OUT / "shortly-devops-slides.pptx"
    prs.save(path)
    print(f"Wrote {path.relative_to(ROOT)} with {len(prs.slides)} slides")
    return 0


if __name__ == "__main__":
    sys.exit(main())
